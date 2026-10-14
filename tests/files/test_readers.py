"""文件格式适配器测试（T-08.2 / FR-13，S4 提示词第十节）。

覆盖 9 种格式：CSV / JSON / JSONL / XLSX / PARQUET / ORC / XML / MARKDOWN / PDF，
外加格式识别（扩展名 + 魔数嗅探）与 Schema 类型推断。
"""
from __future__ import annotations

from pathlib import Path

import pytest

from etl_agent.files.readers import (
    SUPPORTED_FORMATS,
    UnsupportedFormat,
    detect_format,
    infer_columns,
    reader_for,
)
from etl_agent.files.readers.csv_reader import CsvReader
from etl_agent.files.readers.excel_reader import ExcelReader
from etl_agent.files.readers.json_reader import JsonLinesReader, JsonReader
from etl_agent.files.readers.markdown_reader import MarkdownReader
from etl_agent.files.readers.parquet_reader import OrcReader, ParquetReader
from etl_agent.files.readers.pdf_reader import PdfReader
from etl_agent.files.readers.xml_reader import XmlReader

CSV_TEXT = (
    "order_id,customer_email,amount,created_at\n"
    "1,alice@example.com,10.50,2026-10-14\n"
    "2,bob@example.com,,2026-10-15\n"
)

EXPECTED_CSV_ROWS = [
    {
        "order_id": "1",
        "customer_email": "alice@example.com",
        "amount": "10.50",
        "created_at": "2026-10-14",
    },
    {
        "order_id": "2",
        "customer_email": "bob@example.com",
        "amount": None,  # 空单元格归一为 None
        "created_at": "2026-10-15",
    },
]


def test_csv_reader(tmp_path: Path) -> None:
    """CSV 读取正确：表头成键、空单元格归一为 None、类型按可解析性还原（FR-13）。"""
    path = tmp_path / "orders.csv"
    path.write_text(CSV_TEXT, encoding="utf-8")

    rows = list(CsvReader().iter_rows(path))

    assert rows == EXPECTED_CSV_ROWS
    types = {column.name: column.type for column in infer_columns(rows)}
    assert types["order_id"] == "integer"
    assert types["amount"] == "number"
    assert types["created_at"] == "date"
    assert types["customer_email"] == "string"


def test_json_reader(tmp_path: Path) -> None:
    """JSON（顶层数组）读取正确（FR-13）。"""
    path = tmp_path / "orders.json"
    path.write_text(
        '[{"order_id": 1, "amount": 10.5}, {"order_id": 2, "amount": 20.0}]', encoding="utf-8"
    )

    rows = list(JsonReader().iter_rows(path))

    assert rows == [{"order_id": 1, "amount": 10.5}, {"order_id": 2, "amount": 20.0}]
    types = {column.name: column.type for column in infer_columns(rows)}
    assert types["order_id"] == "integer"
    assert types["amount"] == "number"


def test_jsonl_reader(tmp_path: Path) -> None:
    """JSONL（一行一对象）读取正确（FR-13 扩展格式）。"""
    path = tmp_path / "orders.jsonl"
    path.write_text(
        '{"order_id": 1, "ok": true}\n\n{"order_id": 2, "ok": false}\n', encoding="utf-8"
    )

    rows = list(JsonLinesReader().iter_rows(path))

    assert rows == [{"order_id": 1, "ok": True}, {"order_id": 2, "ok": False}]
    types = {column.name: column.type for column in infer_columns(rows)}
    assert types["ok"] == "boolean"


def test_excel_reader(tmp_path: Path) -> None:
    """XLSX 读取正确（首行表头、首工作表，FR-13 点名格式）。"""
    from openpyxl import Workbook

    path = tmp_path / "orders.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["order_id", "customer_email"])
    sheet.append([1, "alice@example.com"])
    sheet.append([2, None])
    workbook.save(path)

    rows = list(ExcelReader().iter_rows(path))

    assert rows == [
        {"order_id": 1, "customer_email": "alice@example.com"},
        {"order_id": 2, "customer_email": None},
    ]


def test_parquet_reader(tmp_path: Path) -> None:
    """Parquet 读取正确（FR-13 点名格式）。"""
    import pyarrow as pa
    import pyarrow.parquet as pq

    path = tmp_path / "orders.parquet"
    pq.write_table(pa.table({"order_id": [1, 2], "amount": [10.5, 20.0]}), path)

    rows = list(ParquetReader().iter_rows(path))

    assert rows == [{"order_id": 1, "amount": 10.5}, {"order_id": 2, "amount": 20.0}]


def test_orc_reader(tmp_path: Path) -> None:
    """ORC 读取正确（FR-13 扩展格式）。"""
    import pyarrow as pa
    from pyarrow import orc

    path = tmp_path / "orders.orc"
    orc.write_table(pa.table({"order_id": [1, 2]}), path)

    rows = list(OrcReader().iter_rows(path))

    assert rows == [{"order_id": 1}, {"order_id": 2}]


def test_xml_reader(tmp_path: Path) -> None:
    """XML 读取正确：根元素直接子元素为记录，属性以 @ 前缀入列（FR-13 扩展格式）。"""
    path = tmp_path / "orders.xml"
    path.write_text(
        '<?xml version="1.0"?>\n<orders>'
        '<order id="1"><amount>10.5</amount></order>'
        '<order id="2"><amount>20.0</amount></order>'
        "</orders>",
        encoding="utf-8",
    )

    rows = list(XmlReader().iter_rows(path))

    assert rows == [
        {"@id": "1", "amount": "10.5"},
        {"@id": "2", "amount": "20.0"},
    ]


def test_markdown_reader(tmp_path: Path) -> None:
    """Markdown 表格读取正确（FR-13 扩展格式）。"""
    path = tmp_path / "orders.md"
    path.write_text(
        "# 订单\n\n"
        "| order_id | customer_email |\n"
        "| --- | --- |\n"
        "| 1 | alice@example.com |\n"
        "| 2 | bob@example.com |\n",
        encoding="utf-8",
    )

    rows = list(MarkdownReader().iter_rows(path))

    assert rows == [
        {"order_id": "1", "customer_email": "alice@example.com"},
        {"order_id": "2", "customer_email": "bob@example.com"},
    ]


def test_markdown_reader_without_table(tmp_path: Path) -> None:
    """Markdown 无表格时退化为一列文本（避免探查结果为空）。"""
    path = tmp_path / "notes.md"
    path.write_text("第一行\n\n第二行\n", encoding="utf-8")

    rows = list(MarkdownReader().iter_rows(path))

    assert rows == [{"text": "第一行"}, {"text": "第二行"}]


def test_pdf_reader(tmp_path: Path) -> None:
    """PDF 文本层抽取：一页一行（FR-13 扩展格式）。"""
    path = tmp_path / "report.pdf"
    path.write_bytes(_minimal_pdf("hello pdf"))

    rows = list(PdfReader().iter_rows(path))

    assert len(rows) == 1
    assert rows[0]["page"] == 1
    assert "hello pdf" in str(rows[0]["text"])


def test_reader_registry_covers_nine_formats() -> None:
    """注册表覆盖全部 9 种格式，且都能取到适配器（FR-13 / FR-64）。"""
    assert len(SUPPORTED_FORMATS) == 9
    for file_format in SUPPORTED_FORMATS:
        assert reader_for(file_format) is not None


def test_detect_format_by_extension() -> None:
    """扩展名通道：9 种格式的常见后缀均识别正确。"""
    assert detect_format("a.csv", b"") == "CSV"
    assert detect_format("a.json", b"") == "JSON"
    assert detect_format("a.jsonl", b"") == "JSONL"
    assert detect_format("a.xlsx", b"") == "XLSX"
    assert detect_format("a.parquet", b"") == "PARQUET"
    assert detect_format("a.orc", b"") == "ORC"
    assert detect_format("a.xml", b"") == "XML"
    assert detect_format("a.md", b"") == "MARKDOWN"
    assert detect_format("a.pdf", b"") == "PDF"


def test_detect_format_by_magic() -> None:
    """魔数通道优先于扩展名：内容权威（S4 提示词 9.3）。"""
    assert detect_format("mystery.bin", b"%PDF-1.4 ...") == "PDF"
    assert detect_format("mystery.bin", b"PAR1....") == "PARQUET"
    assert detect_format("mystery.bin", b"ORC....") == "ORC"
    assert detect_format("mystery.bin", b"PK\x03\x04....") == "XLSX"
    assert detect_format("mystery.bin", b'<?xml version="1.0"?>') == "XML"


def test_detect_format_unsupported() -> None:
    """两通道都认不出 → 拒绝，不猜格式硬读（FR-13）。"""
    with pytest.raises(UnsupportedFormat):
        detect_format("notes.txt", b"just plain text")
    with pytest.raises(UnsupportedFormat):
        reader_for("AVRO")


def _minimal_pdf(text: str) -> bytes:
    """构造最小可用单页 PDF（用于 PDF 适配器测试，不引入额外依赖）。"""
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
        ),
        b"",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode("ascii")
    objects[3] = (
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
    )

    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_offset = len(out)
    size = len(objects) + 1
    out += f"xref\n0 {size}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {size} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode()
    return bytes(out)
