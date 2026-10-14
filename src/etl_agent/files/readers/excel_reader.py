"""M-E 文件资产 — XLSX 适配器。

SPEC 依据：etl_agent-SPEC_v1.0.3-Baseline 第 3.2 章 M-E
覆盖 FR：FR-13（Excel/XLSX 为需求文档点名格式）、FR-64（点名格式覆盖核对）
对应 AC：AC-02（只产出内存行，脱敏由 masking 统一负责）

只读首个工作表：需求文档点名的场景是「单表文件资产」，多 Sheet 的语义
（选哪个 Sheet）不在 S4 范围内，静默合并会制造错误的 Schema。
"""
from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path

from etl_agent.files.readers.base import ReaderError, Row


class ExcelReader:
    """XLSX 适配器（read_only 模式，首行表头，惰性逐行）。"""

    def iter_rows(self, path: Path) -> Iterator[Row]:
        """按行迭代首个工作表（FR-13）。

        Raises:
            ReaderError: 工作簿无法打开或没有工作表。
        """
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException

        try:
            # read_only=True：openpyxl 以流式方式读取，不把整份工作簿建树到内存
            workbook = load_workbook(filename=path, read_only=True, data_only=True)
        except (InvalidFileException, OSError, ValueError) as exc:
            raise ReaderError("XLSX 文件无法解析") from exc
        try:
            worksheet = workbook.worksheets[0] if workbook.worksheets else None
            if worksheet is None:
                raise ReaderError("XLSX 不含任何工作表")
            rows = worksheet.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                raise ReaderError("XLSX 缺少表头行")
            names = [_cell_name(index, value) for index, value in enumerate(header)]
            for values in rows:
                yield {name: _value_at(values, index) for index, name in enumerate(names)}
        finally:
            workbook.close()


def _cell_name(index: int, value: object) -> str:
    """表头单元格 → 列名（空表头用列序号占位，避免出现 None 键）。"""
    text = "" if value is None else str(value).strip()
    return text or f"column_{index + 1}"


def _value_at(values: tuple[object, ...], index: int) -> object:
    """按序号取值（行长度短于表头时补 None）。"""
    return values[index] if index < len(values) else None
